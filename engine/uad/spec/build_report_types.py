"""The Restricted Appraisal Update Report and the Completion Report, as data.

UAD 3.6 has three reports: the URAR (A-1..H-1), the Restricted Appraisal
Update Report (A-2, B-2, C-2, D-2, F-2, H-2) and the Completion Report (A-3,
B-3, C-3, D-3, F-3, H-3).  Each update or completion report is submitted
against an existing URAR and must match its format (UAD FAQ 9).  This reads
each report's own appendices -- nothing typed in -- into
``spec/report_types/<kind>.json``:

  sections     from the Implementation Guide's table of contents
  fields       per Report Field ID, from the Delivery Specification: section,
               subsection, label, data points (xPath, format, enumerations with
               their display text, conditionality), as A-1's field_index has
  display      the Implementation Guide's display-rule tables (B-2 / B-3)
  rules        the compliance rules (H-2 / H-3)
  sample       the D-2 / D-3 sample XML(s)
  form         the report as the Layout (C-2 / C-3) draws it, read by the same
               extractor as C-1 (``extract_c1_form.extract_page``) -- the two
               layouts use C-1's own primitives
  furniture    the header and footer labels (HF.001-HF.008), from the
               Delivery Specification's Report Location column

    python -m uad.spec.build_report_types        (from engine/)
"""
from __future__ import annotations

import json
import re
from collections import defaultdict
from pathlib import Path

import openpyxl
import pymupdf

from ..paths import PROJECT_ROOT

# the update and completion appendices live in specs/official (copied from the
# owner's UAD36-Blueprint repo, identical to the GSE downloads)
SPEC_ROOT = PROJECT_ROOT / "specs" / "official"
from . import extract_display_rules as edr
from .build_index import discriminators, parse_locations, requires

OUT = Path(__file__).resolve().parents[2] / "spec" / "report_types"

KINDS = {
    "update": {"title": "Restricted Appraisal Update Report", "folder": "Restricted_Update",
               "layout": "C-2_Layout",
               "delivery": "A-2_Delivery_Spec.xlsx", "guide": "B-2_Implementation_Guide.pdf",
               "rules": "H-2_Compliance_Rules.xlsx", "samples": "*_AppraisalUpdate_v1.*.xml",
               "document_type": "ValuationUpdate"},
    "completion": {"title": "Completion Report", "folder": "Completion_Report",
                   "layout": "C-3_Layout",
                   "delivery": "A-3_Delivery_Spec.xlsx", "guide": "B-3_Implementation_Guide.pdf",
                   "rules": "H-3_Compliance_Rules.xlsx", "samples": "*_CompletionReport_v1.*.xml",
                   "document_type": "ValuationCompletion"},
}
TOC = re.compile(r"^(\d{2})\s+([A-Z][A-Za-z ,/&'-]+?)\s*(?:[_.]{3,}|\s{2,})\s*\d+\s*$")


def _t(v) -> str:
    return " ".join(str(v or "").split())


def _sheet(path: Path, prefix: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[[n for n in wb.sheetnames if n.startswith(prefix) and "Marked" not in n][0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    head_i = next(i for i, r in enumerate(rows) if r and any(c and "Unique ID" in str(c) for c in r))
    head = [_t(c) for c in rows[head_i]]
    return [dict(zip(head, r)) for r in rows[head_i + 1:] if r and any(c is not None for c in r)]


def sections_of(guide: Path) -> dict[int, str]:
    """The numbered chapters, from the guide's table of contents."""
    out: dict[int, str] = {}
    doc = pymupdf.open(str(guide))
    for pno in range(min(6, len(doc))):
        for line in doc[pno].get_text().splitlines():
            m = TOC.match(line.strip())
            if m:
                out.setdefault(int(m.group(1)), m.group(2).strip())
    return out


HF = re.compile(r"^(HF\.\d{3}):\s*Header and Footer\s*-\s*(Header|Footer)\s*-\s*(.*)$")


def furniture(delivery: list[dict]) -> dict[str, dict]:
    """HF.001-HF.008: where each header/footer item sits and what labels it."""
    out: dict[str, dict] = {}
    for row in delivery:
        for part in _t(row.get("Report Location")).split("|"):
            m = HF.match(part.strip())
            if m and m.group(1) not in out:
                out[m.group(1)] = {"place": m.group(2), "label": m.group(3).strip(),
                                   "mismo": _t(row.get("MISMO Data Point Name"))}
    return dict(sorted(out.items()))


def layout(folder: Path, name: str) -> dict:
    """The Layout appendix as C-1's extractor reads it.  The first two pages
    are the cover and the legend, as in C-1."""
    from . import extract_c1_form as c1
    base = next(folder.rglob(name))
    clean = next(p for p in base.glob("*.pdf") if "without" in p.name)
    with_ids = next(p for p in base.glob("*.pdf") if "without" not in p.name)
    a, b = pymupdf.open(str(clean)), pymupdf.open(str(with_ids))
    try:
        pages = [c1.extract_page(a[n], b[n], n + 1) for n in range(2, a.page_count)]
    finally:
        a.close()
        b.close()
    return {"source": [clean.name, with_ids.name], "pages": [_repair(p) for p in pages]}


def _repair(page: dict) -> dict:
    """What C-1's extractor misreads on C-3's Completion Status (p3), fixed by
    shape rather than by name:

    * two Yes/No questions set one under the other came out as a bold head
      (the first question) and one check carrying both field IDs: each
      question becomes its own check with its own ID;
    * the second question wraps, and its last word ("specifications?") was
      read as the first column head of the table below it: a head that starts
      with words ending in "?" gives them back to the check above;
    * a table row whose IDs all sit in one cell, one per column, gets one ID
      per column again.
    """
    els = page["elements"]
    out: list[dict] = []
    for i, e in enumerate(els):
        if (e["kind"] == "check" and len(e.get("ids") or []) >= 2 and out
                and out[-1]["kind"] == "head" and out[-1].get("bold")
                and (out[-1].get("text") or "").rstrip().endswith("?")
                and 0 < e["y"] - out[-1]["y"] < 16):
            head = out.pop()
            # C-3 sets "Yes No" over the pair (p3); merged, the header was lost
            first = {**e, "y": head["y"], "y1": head["y1"], "text": head["text"],
                     "ids": e["ids"][:1], "yes_no": True}
            second = {**e, "ids": e["ids"][1:], "yes_no": False}
            nxt = els[i + 1] if i + 1 < len(els) else None
            if nxt is not None and nxt["kind"] == "table" and nxt.get("columns"):
                col = nxt["columns"][0]
                m = re.match(r"^(.*?\?)\s+(.*)$", col.get("head", ""))
                if m:
                    second["text"] = f"{second['text']} {m.group(1)}"
                    col["head"] = m.group(2)
            out += [first, second]
            continue
        if e["kind"] == "table":
            n = len(e.get("columns") or [])
            for row in e.get("rows") or []:
                cells = row.get("cells") or []
                full = [c for c in cells if c.get("ids")]
                if n > 1 and len(cells) == n and len(full) == 1 and len(full[0]["ids"]) == n:
                    ids = full[0]["ids"]
                    for c, fid in zip(cells, ids):
                        c["ids"] = [fid]
        out.append(e)
    page["elements"] = out
    return page


def build(kind: str) -> dict:
    cfg = KINDS[kind]
    folder = next(SPEC_ROOT.rglob(cfg["folder"]))
    guide = next(folder.rglob(cfg["guide"]))
    sections = sections_of(guide)

    fields: dict[str, dict] = {}
    enums_by_uid: dict[str, list[dict]] = defaultdict(list)
    for row in _sheet(next(folder.rglob(cfg["delivery"])), "UAD Enumerations"):
        uid = _t(row.get("Unique ID"))
        value = _t(row.get("Supported Enumerations"))
        display_col = next((k for k in row if k.startswith("Output Report Display Text")), "")
        display = _t(str(row.get(display_col) or "").split("\n")[0])
        if uid and value:
            enums_by_uid[uid].append({"value": value,
                                      "display": "" if display == "<<Does Not Display>>" else display})
    not_shown = []
    delivery = _sheet(next(folder.rglob(cfg["delivery"])), "UAD Delivery Spec")
    for row in delivery:
        location = _t(row.get("Report Location"))
        dp = {"unique_id": _t(row.get("Unique ID")), "context": _t(row.get("Additional Context")),
              "mismo": _t(row.get("MISMO Data Point Name")), "attribute": _t(row.get("MISMO Attribute Name")),
              "xpath": _t(row.get("MISMO xPath")), "container": _t(row.get("MISMO Parent Container")),
              "format_type": _t(row.get("UAD Data Point Format Type")),
              "format_details": _t(row.get("UAD Data Point Format Details")),
              "conditionality": _t(row.get("Conditionality")),
              "conditionality_details": _t(row.get("Conditionality Details")),
              "requires": requires(_t(row.get("Conditionality Details"))),
              "definition": _t(row.get(next((k for k in row if k.startswith("MISMO Data Point / Attribute")), "")))}
        if enums_by_uid.get(dp["unique_id"]):
            dp["enumerations"] = enums_by_uid[dp["unique_id"]]
        places = parse_locations(location)
        if not places:
            not_shown.append(dp)
            continue
        for fid, section, subsection, label in places:
            entry = fields.setdefault(fid, {"field_id": fid, "section": section, "subsection": subsection,
                                            "label": label, "data_points": [], "display_rules": []})
            entry["data_points"].append(dict(dp))

    rules = edr.consolidate(edr.extract(guide, sections, carry=True))
    for fid, rec in rules.items():
        fid = re.sub(r"^(\d)\.", r"0\1.", fid)       # B-3 prints "8.010" for 08.010
        if fid not in fields:
            # A field the guide rules on but A-2/A-3 gives no data point: the
            # margin slot of a certification paragraph, whose rule decides
            # whether the paragraph prints (B-2 05.011 "Always displays").
            st = rec.get("statements") or [{}]
            fields[fid] = {"field_id": fid, "section": st[0].get("section", ""), "subsection": "",
                           "label": "", "data_points": [], "display_rules": []}
        fields[fid]["display_rules"] = fields[fid]["display_rules"] + rec.get("statements", [])

    compliance = []
    for row in _sheet(next(folder.rglob(cfg["rules"])), "UAD Compliance Rules"):
        if _t(row.get("Message ID")):
            compliance.append({k: _t(v) for k, v in row.items() if k})

    samples = sorted(str(p.relative_to(PROJECT_ROOT)) for p in folder.rglob(cfg["samples"]))
    return {"kind": kind, "title": cfg["title"], "document_type": cfg["document_type"],
            "sources": {"delivery": cfg["delivery"], "guide": cfg["guide"], "rules": cfg["rules"]},
            "sections": {str(k): v for k, v in sorted(sections.items())},
            "fields": dict(sorted(fields.items(), key=lambda kv: [int(x) for x in re.findall(r"\d+", kv[0])])),
            "not_on_report": not_shown,
            # per Additional Context, the type value that names an instance --
            # "Appraiser" -> PartyRoleType Appraiser -- as A-1's field_index has
            "discriminators": discriminators(delivery),
            "furniture": furniture(delivery),
            "form": layout(folder, cfg["layout"]),
            "compliance_rules": compliance, "samples": samples}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for kind in KINDS:
        data = build(kind)
        path = OUT / f"{kind}.json"
        path.write_text(json.dumps(data, indent=1, ensure_ascii=False), encoding="utf-8")
        ruled = sum(1 for f in data["fields"].values() if f["display_rules"])
        print(f"{kind:11s} sections={data['sections']}")
        print(f"            fields={len(data['fields'])} with display rules={ruled} "
              f"not-on-report points={len(data['not_on_report'])} rules={len(data['compliance_rules'])} "
              f"samples={data['samples']}")


if __name__ == "__main__":
    main()
